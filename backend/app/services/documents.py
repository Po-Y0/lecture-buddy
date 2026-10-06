from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from fastapi import HTTPException, UploadFile
from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from starlette.concurrency import run_in_threadpool

from backend.app.config import get_settings
from backend.app.models import new_id

SUPPORTED_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    },
    ".pptx": {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    },
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain"},
}
MAX_UNCOMPRESSED_ARCHIVE_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2_000
MAX_EXTRACTED_TEXT_CHARACTERS = 5_000_000


def safe_filename(filename: str | None) -> str:
    normalized = (filename or "").replace("\\", "/")
    name = Path(normalized).name.strip()
    if (
        not name
        or len(name) > 255
        or any(ord(char) < 32 or ord(char) == 127 for char in name)
    ):
        raise HTTPException(status_code=400, detail="Invalid file name")
    return name


def _validate_type(name: str, mime_type: str | None, payload: bytes) -> tuple[str, str]:
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_TYPES:
        raise HTTPException(status_code=415, detail="Unsupported file type")
    expected = SUPPORTED_TYPES[extension]
    actual = (mime_type or "").lower()
    if actual and actual not in expected and actual != "application/octet-stream":
        raise HTTPException(
            status_code=415, detail="File type does not match its extension"
        )
    if extension == ".pdf" and not payload.lstrip().startswith(b"%PDF-"):
        raise HTTPException(status_code=415, detail="File is not a valid PDF")
    if extension in {".docx", ".pptx"} and not payload.startswith(b"PK"):
        raise HTTPException(
            status_code=415, detail="Office document is not a valid ZIP package"
        )
    if extension in {".txt", ".md"} and b"\x00" in payload:
        raise HTTPException(status_code=415, detail="Text file contains binary data")
    return extension, actual or next(iter(expected))


def _validate_archive(payload: bytes) -> None:
    try:
        with ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise HTTPException(
                    status_code=413, detail="Office document contains too many entries"
                )
            if (
                sum(entry.file_size for entry in entries)
                > MAX_UNCOMPRESSED_ARCHIVE_BYTES
            ):
                raise HTTPException(
                    status_code=413,
                    detail="Office document expands beyond the processing limit",
                )
    except BadZipFile as error:
        raise HTTPException(
            status_code=415, detail="Office document is not a valid package"
        ) from error


def extract_text(extension: str, payload: bytes) -> str:
    try:
        if extension in {".txt", ".md"}:
            text = payload.decode("utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(payload), strict=True)
            if reader.is_encrypted:
                raise HTTPException(
                    status_code=422, detail="Password-protected PDFs are not supported"
                )
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            document = Document(BytesIO(payload))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            text = "\n".join(parts)
        elif extension == ".pptx":
            presentation = Presentation(BytesIO(payload))
            text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            )
        else:
            raise HTTPException(status_code=415, detail="Unsupported file type")
    except UnicodeDecodeError as error:
        raise HTTPException(
            status_code=422, detail="Text files must use UTF-8 encoding"
        ) from error
    except HTTPException:
        raise
    except (
        PdfReadError,
        DocxPackageNotFoundError,
        PptxPackageNotFoundError,
        XMLSyntaxError,
        ValueError,
        OSError,
        KeyError,
    ) as error:
        raise HTTPException(
            status_code=422, detail="Document could not be parsed"
        ) from error

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) > MAX_EXTRACTED_TEXT_CHARACTERS:
        raise HTTPException(
            status_code=413,
            detail="Extracted document text exceeds the processing limit",
        )
    if not text:
        detail = (
            "No selectable text found. OCR for scanned or image-only PDFs is not supported."
            if extension == ".pdf"
            else "No readable text was found in the document"
        )
        raise HTTPException(status_code=422, detail=detail)
    return text


async def read_and_store(upload: UploadFile, owner_id: str) -> dict:
    name = safe_filename(upload.filename)
    settings = get_settings()
    payload = await upload.read(settings.max_upload_bytes + 1)
    if not payload:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    if len(payload) > settings.max_upload_bytes:
        raise HTTPException(
            status_code=413, detail="File exceeds the configured upload limit"
        )

    extension, mime_type = _validate_type(name, upload.content_type, payload)
    if extension in {".docx", ".pptx"}:
        _validate_archive(payload)
    extracted_text = await run_in_threadpool(extract_text, extension, payload)

    source_id = new_id()
    owner_dir = Path(settings.upload_dir).resolve() / owner_id
    owner_dir.mkdir(parents=True, exist_ok=True)
    storage_path = owner_dir / f"{source_id}{extension}"
    await run_in_threadpool(storage_path.write_bytes, payload)
    return {
        "id": source_id,
        "name": name,
        "mime_type": mime_type,
        "size": len(payload),
        "storage_path": str(storage_path),
        "extraction_status": "ready",
        "extracted_text": extracted_text,
    }

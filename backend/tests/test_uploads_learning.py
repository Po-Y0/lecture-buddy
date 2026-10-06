import asyncio
import json
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from backend.app.config import get_settings
from backend.app.services.ai import AIServiceError, generate_questions
from backend.app.services.documents import extract_text


def register_and_login(client: TestClient):
    client.post(
        "/api/auth/register",
        json={
            "name": "Student",
            "email": "student@example.com",
            "password": "safePassword1",
        },
    )
    return client.post(
        "/api/auth/login",
        json={"email": "student@example.com", "password": "safePassword1"},
    )


def test_extracts_text_from_supported_office_formats():
    word = Document()
    word.add_paragraph("Word lecture content")
    word_bytes = BytesIO()
    word.save(word_bytes)
    assert "Word lecture content" in extract_text(".docx", word_bytes.getvalue())

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Presentation lecture content"
    deck_bytes = BytesIO()
    deck.save(deck_bytes)
    assert "Presentation lecture content" in extract_text(
        ".pptx", deck_bytes.getvalue()
    )


def test_upload_download_and_delete_source(client, monkeypatch, tmp_path):
    register_and_login(client)
    settings = get_settings().model_copy(update={"upload_dir": str(tmp_path)})
    monkeypatch.setattr("backend.app.services.documents.get_settings", lambda: settings)
    course = client.post("/api/courses", json={"name": "Biology"}).json()
    response = client.post(
        f"/api/courses/{course['id']}/sources",
        files={"file": ("notes.txt", b"Cell structures and function", "text/plain")},
    )
    assert response.status_code == 201
    source = response.json()
    assert source["name"] == "notes.txt"
    assert source["size"] == len(b"Cell structures and function")
    assert source["extractionStatus"] == "ready"

    stored_file = tmp_path / client.get("/api/sources").json()[0]["id"]
    actual_path = next(tmp_path.rglob("*.txt"))
    assert actual_path.is_file()
    download = client.get(f"/api/sources/{source['id']}/download")
    assert download.status_code == 200
    assert download.content == b"Cell structures and function"
    assert client.delete(f"/api/sources/{source['id']}").status_code == 204
    assert not actual_path.exists()
    assert client.get("/api/sources").json() == []
    assert not stored_file.exists()


def test_upload_rejects_unsupported_mismatched_and_scanned_files(client):
    register_and_login(client)
    course = client.post("/api/courses", json={"name": "Physics"}).json()
    unsupported = client.post(
        f"/api/courses/{course['id']}/sources",
        files={
            "file": ("payload.exe", b"not an allowed file", "application/octet-stream")
        },
    )
    assert unsupported.status_code == 415

    mismatch = client.post(
        f"/api/courses/{course['id']}/sources",
        files={"file": ("lecture.pdf", b"plain text", "application/pdf")},
    )
    assert mismatch.status_code == 415

    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    scanned_like = BytesIO()
    writer.write(scanned_like)
    response = client.post(
        f"/api/courses/{course['id']}/sources",
        files={"file": ("scan.pdf", scanned_like.getvalue(), "application/pdf")},
    )
    assert response.status_code == 422
    assert "OCR" in response.json()["detail"]


def test_upload_enforces_configured_size_limit(client, monkeypatch, tmp_path):
    register_and_login(client)
    settings = get_settings().model_copy(
        update={"upload_dir": str(tmp_path), "max_upload_bytes": 1024}
    )
    monkeypatch.setattr("backend.app.services.documents.get_settings", lambda: settings)
    course = client.post("/api/courses", json={"name": "Chemistry"}).json()
    response = client.post(
        f"/api/courses/{course['id']}/sources",
        files={"file": ("lecture.txt", b"x" * 1025, "text/plain")},
    )
    assert response.status_code == 413


def test_source_access_is_private_to_owner(client, monkeypatch, tmp_path):
    register_and_login(client)
    settings = get_settings().model_copy(update={"upload_dir": str(tmp_path)})
    monkeypatch.setattr("backend.app.services.documents.get_settings", lambda: settings)
    course = client.post("/api/courses", json={"name": "Literature"}).json()
    uploaded = client.post(
        f"/api/courses/{course['id']}/sources",
        files={"file": ("reading.md", b"Study notes", "text/markdown")},
    ).json()

    client.post(
        "/api/auth/register",
        json={
            "name": "Other",
            "email": "other@example.com",
            "password": "safePassword2",
        },
    )
    client.cookies.clear()
    client.post(
        "/api/auth/login",
        json={"email": "other@example.com", "password": "safePassword2"},
    )
    assert client.get("/api/sources").json() == []
    assert client.get(f"/api/sources/{uploaded['id']}/download").status_code == 404
    assert client.delete(f"/api/sources/{uploaded['id']}").status_code == 404


def test_ai_routes_report_missing_provider_credentials(client, monkeypatch):
    register_and_login(client)
    course = client.post("/api/courses", json={"name": "History"}).json()
    client.post(
        f"/api/courses/{course['id']}/notes",
        json={"title": "Lecture", "content": "Readable course material."},
    )
    settings = get_settings().model_copy(update={"openai_api_key": None})
    monkeypatch.setattr("backend.app.services.ai.get_settings", lambda: settings)

    response = client.post(f"/api/courses/{course['id']}/summarize")
    assert response.status_code == 503
    assert "OPENAI_API_KEY" in response.json()["detail"]


def test_ai_quiz_response_parser_validates_generated_output(monkeypatch):
    questions = [
        {
            "question": f"Question {index}",
            "choices": ["One", "Two", "Three", "Four"],
            "answerIndex": index % 4,
            "explanation": "Because the source says so.",
        }
        for index in range(10)
    ]

    async def valid_response(_system, _user):
        return json.dumps(questions)

    monkeypatch.setattr("backend.app.services.ai._complete", valid_response)
    parsed = asyncio.run(generate_questions("Course material", 10))
    assert len(parsed) == 10
    assert parsed[0]["answerIndex"] == 0
    assert parsed[0]["id"]

    async def invalid_response(_system, _user):
        return "[]"

    monkeypatch.setattr("backend.app.services.ai._complete", invalid_response)
    try:
        asyncio.run(generate_questions("Course material", 10))
    except AIServiceError as error:
        assert error.status_code == 502
        assert "invalid quiz" in error.detail
    else:
        raise AssertionError("Invalid quiz output must be rejected")


def test_ai_summary_and_quiz_attempts_are_persisted_and_scored(client, monkeypatch):
    register_and_login(client)
    course = client.post("/api/courses", json={"name": "Computer Science"}).json()
    client.post(
        f"/api/courses/{course['id']}/notes",
        json={"title": "Trees", "content": "A binary tree has at most two children."},
    )

    settings = get_settings().model_copy(update={"openai_api_key": "test-key"})
    monkeypatch.setattr("backend.app.services.ai.get_settings", lambda: settings)

    async def mocked_summarize(_material):
        return "Binary trees have up to two children."

    monkeypatch.setattr(
        "backend.app.routers.learning.summarize",
        mocked_summarize,
    )
    summary = client.post(f"/api/courses/{course['id']}/summarize")
    assert summary.status_code == 200
    assert summary.json()["summary"].startswith("Binary trees")

    questions = [
        {
            "id": f"q-{index}",
            "question": f"Question {index}",
            "choices": ["A", "B", "C", "D"],
            "answerIndex": index % 4,
            "explanation": "Explanation",
        }
        for index in range(10)
    ]

    async def mocked_generate(_material, count):
        assert count == 10
        return questions

    monkeypatch.setattr(
        "backend.app.routers.learning.generate_questions", mocked_generate
    )
    generated = client.post(
        f"/api/courses/{course['id']}/quiz",
        json={"count": 10},
    )
    assert generated.status_code == 201
    quiz = generated.json()
    assert len(quiz["questions"]) == 10

    attempt = client.post(
        f"/api/quizzes/{quiz['quizId']}/attempts",
        json={
            "answers": [
                {"questionId": "q-0", "answerIndex": 0},
                {"questionId": "q-1", "answerIndex": 0},
            ]
        },
    )
    assert attempt.status_code == 201
    assert attempt.json()["score"] == 1
    assert attempt.json()["total"] == 2
    assert client.get("/api/quiz-attempts").json() == [attempt.json()]

    duplicate_answers = client.post(
        f"/api/quizzes/{quiz['quizId']}/attempts",
        json={
            "answers": [
                {"questionId": "q-0", "answerIndex": 0},
                {"questionId": "q-0", "answerIndex": 1},
            ]
        },
    )
    assert duplicate_answers.status_code == 422
